"""Generate a reproducible small synthetic corpus for weakly supervised NER and ranking."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "synthetic" / "bootstrapped_resumes"

FIRST_NAMES = [
    "Asha", "Daniel", "Leila", "Mateo", "Nora", "Priya", "Jonas", "Mina", "Owen", "Rina",
    "Kai", "Sofia", "Ethan", "Amara", "Luca", "Ivy", "Noah", "Zara", "Milo", "Anika",
]
LAST_NAMES = [
    "Mehta", "Brooks", "Haddad", "Silva", "Chen", "Patel", "Fischer", "Okafor", "Nguyen", "Rossi",
]
COMPANIES = ["Northstar Labs", "Blue Oak Systems", "Cedar Health", "Orbit Commerce", "Maple Analytics"]
LOCATIONS = ["Austin, TX", "Boston, MA", "Chicago, IL", "Denver, CO", "Seattle, WA"]
BACKEND_SKILLS = ["Python", "FastAPI", "PostgreSQL", "Docker", "REST APIs", "SQL", "Git"]
DATA_SKILLS = ["Python", "Pandas", "SQL", "Tableau", "Git"]
FRONTEND_SKILLS = ["JavaScript", "TypeScript", "React", "CSS", "Accessibility", "Git"]


def add_paragraphs(document: Document, lines: list[str]) -> None:
    for line in lines:
        document.add_paragraph(line)


def build_resume(index: int) -> tuple[Path, dict[str, object]]:
    first = FIRST_NAMES[index % len(FIRST_NAMES)]
    last = LAST_NAMES[(index * 3) % len(LAST_NAMES)]
    name = f"{first} {last} {index + 1:02d}"
    location = LOCATIONS[index % len(LOCATIONS)]
    company = COMPANIES[index % len(COMPANIES)]
    start_year = 2017 + (index % 6)
    end_year = start_year + 3 + (index % 3)
    target = index % 2 == 0
    if target:
        title = ["Backend Engineer", "Platform Engineer", "API Engineer"][index % 3]
        skills = BACKEND_SKILLS[: 4 + (index % 4)]
        summary = f"Backend engineer building reliable Python services and APIs for {company}."
    elif index % 3 == 0:
        title = "Data Analyst"
        skills = DATA_SKILLS[: 3 + (index % 3)]
        summary = f"Data analyst turning operational data into decisions at {company}."
    else:
        title = "Frontend Developer"
        skills = FRONTEND_SKILLS[: 3 + (index % 4)]
        summary = f"Frontend developer delivering accessible web experiences for {company}."

    document = Document()
    document.add_paragraph("SYNTHETIC RESUME")
    document.add_paragraph(f"{name} | {name.lower().replace(' ', '.')}@example.com | +1 555 010 {index:04d} | {location}")
    document.add_paragraph("SUMMARY")
    document.add_paragraph(summary)
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph(f"{title}, {company} | {start_year} - {end_year}")
    document.add_paragraph(f"Delivered production systems using {' and '.join(skills[:3])}.")
    document.add_paragraph(f"Collaborated with product teams at {COMPANIES[(index + 1) % len(COMPANIES)]}.")
    document.add_paragraph("EDUCATION")
    document.add_paragraph(f"B.S. Computer Science, {['State University', 'Metro Institute', 'Lakeside College'][index % 3]} | {start_year - 1}")
    document.add_paragraph("SKILLS")
    document.add_paragraph(", ".join(skills))
    if index % 4 == 0:
        document.add_paragraph(f"Additional project: migrated services from {COMPANIES[(index + 2) % len(COMPANIES)]}.")

    path = OUTPUT / f"resume_{index + 1:03d}_{first.lower()}_{last.lower()}.docx"
    document.save(path)
    benchmark_relevant = target and index in {0, 8, 16, 24, 32, 40}
    return path, {
        "name": name,
        "target_backend": target,
        "benchmark_relevant": benchmark_relevant,
        "skills": skills,
        "title": title,
    }


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for old in OUTPUT.glob("resume_*.docx"):
        old.unlink()
    manifest = []
    for index in range(50):
        _, metadata = build_resume(index)
        manifest.append(metadata)
    (OUTPUT / "manifest.json").write_text(json.dumps({"data_status": "synthetic_generated", "records": manifest}, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": str(OUTPUT),
        "documents": len(manifest),
        "target_backend": sum(item["target_backend"] for item in manifest),
        "benchmark_relevant": sum(item["benchmark_relevant"] for item in manifest),
    }, indent=2))


if __name__ == "__main__":
    main()
