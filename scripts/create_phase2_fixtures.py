"""Create labeled synthetic resume fixtures for Phase 2 verification."""

from __future__ import annotations

from pathlib import Path

from docx import Document


RESUMES = (
    (
        "resume_01_asha.docx",
        "Asha Mehta | asha.synthetic@example.invalid | +1-555-0101 | Bengaluru, India",
        "Backend engineer with five years building reliable APIs and data services.",
        "Software Engineer, Northstar Labs | 2021-2026 | Built Python and FastAPI services; maintained PostgreSQL pipelines.",
        "B.Tech in Computer Science, National Institute of Technology | 2020",
        "Python, FastAPI, PostgreSQL, Docker, REST APIs",
    ),
    (
        "resume_02_daniel.docx",
        "Daniel Brooks | daniel.synthetic@example.invalid | +1-555-0102 | Toronto, Canada",
        "Data analyst with four years of experience turning product data into reports and experiments.",
        "Data Analyst, Cedar Analytics | 2022-2026 | Wrote SQL models and Python notebooks; delivered product dashboards.",
        "B.Sc. in Statistics, Western University | 2022",
        "SQL, Python, pandas, Tableau, A/B testing",
    ),
    (
        "resume_03_leila.docx",
        "Leila Haddad | leila.synthetic@example.invalid | +1-555-0103 | Lyon, France",
        "Machine learning engineer with three years deploying classification models.",
        "ML Engineer, Orbit Systems | 2023-2026 | Trained scikit-learn models and deployed batch inference jobs with Docker.",
        "M.Sc. in Data Science, Universite de Lyon | 2023",
        "Python, scikit-learn, Docker, model evaluation, Git",
    ),
    (
        "resume_04_mateo.docx",
        "Mateo Silva | mateo.synthetic@example.invalid | +1-555-0104 | Sao Paulo, Brazil",
        "Frontend developer with six years creating accessible web applications.",
        "Frontend Developer, Vista Digital | 2020-2026 | Built React interfaces and tested reusable TypeScript components.",
        "B.A. in Design and Technology, Federal University of Parana | 2019",
        "JavaScript, TypeScript, React, CSS, accessibility",
    ),
)


def write_resume(path: Path, contact: str, summary: str, experience: str, education: str, skills: str) -> None:
    document = Document()
    document.add_paragraph("SYNTHETIC RESUME - FOR TESTING ONLY")
    document.add_paragraph(contact)
    document.add_paragraph("SUMMARY")
    document.add_paragraph(summary)
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph(experience)
    document.add_paragraph("EDUCATION")
    document.add_paragraph(education)
    document.add_paragraph("SKILLS")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Technical skills"
    table.rows[0].cells[1].text = skills
    document.save(path)


def write_sparse_pdf(path: Path) -> None:
    pdf_stream = b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length 0 >>\nstream\n\nendstream",
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode())
        content.extend(obj)
        content.extend(b"\nendobj\n")
    xref_offset = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode())
    content.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    path.write_bytes(content)


def main() -> None:
    fixture_root = Path(__file__).resolve().parents[1] / "data" / "synthetic" / "resumes"
    fixture_root.mkdir(parents=True, exist_ok=True)
    for filename, contact, summary, experience, education, skills in RESUMES:
        write_resume(fixture_root / filename, contact, summary, experience, education, skills)
    write_sparse_pdf(fixture_root / "scanned_like.pdf")
    print(f"Created {len(RESUMES)} synthetic DOCX resumes and one sparse PDF in {fixture_root}")


if __name__ == "__main__":
    main()
