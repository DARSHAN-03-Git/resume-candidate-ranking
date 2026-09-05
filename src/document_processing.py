"""PDF and DOCX text extraction for Phase 2.

The extractor returns source metadata and quality signals so later stages can
refuse to treat empty or likely-scanned text as trustworthy evidence.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Literal

from docx import Document
from pdfminer.high_level import extract_pages, extract_text


FileType = Literal["pdf", "docx"]
SUPPORTED_SUFFIXES = {".pdf": "pdf", ".docx": "docx"}
SCANNED_TEXT_THRESHOLD = 40


@dataclass(frozen=True)
class ExtractionResult:
    source_path: str
    file_type: FileType
    text: str
    page_count: int
    paragraph_count: int
    table_count: int
    character_count: int
    word_count: int
    likely_scanned: bool
    warnings: tuple[str, ...]


def extract_document(path: str | Path) -> ExtractionResult:
    """Extract text from a PDF or DOCX using a stable result schema."""
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Document does not exist: {source}")

    file_type = SUPPORTED_SUFFIXES.get(source.suffix.lower())
    if file_type is None:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ValueError(f"Unsupported document type {source.suffix!r}; use {supported}")

    if file_type == "pdf":
        text, page_count, paragraph_count, table_count = _extract_pdf(source)
    else:
        text, page_count, paragraph_count, table_count = _extract_docx(source)

    normalized_text = _normalize_text(text)
    character_count = len(normalized_text)
    word_count = len(normalized_text.split())
    likely_scanned = (
        file_type == "pdf"
        and page_count > 0
        and character_count < SCANNED_TEXT_THRESHOLD
    )
    warnings: list[str] = []
    if not normalized_text:
        warnings.append("No extractable text was found.")
    if likely_scanned:
        warnings.append(
            "PDF appears image-only or nearly empty; OCR was not run automatically."
        )

    return ExtractionResult(
        source_path=str(source),
        file_type=file_type,
        text=normalized_text,
        page_count=page_count,
        paragraph_count=paragraph_count,
        table_count=table_count,
        character_count=character_count,
        word_count=word_count,
        likely_scanned=likely_scanned,
        warnings=tuple(warnings),
    )


def _extract_pdf(path: Path) -> tuple[str, int, int, int]:
    page_count = sum(1 for _ in extract_pages(str(path)))
    text = extract_text(str(path)) or ""
    paragraph_count = sum(bool(line.strip()) for line in text.splitlines())
    return text, page_count, paragraph_count, 0


def _extract_docx(path: Path) -> tuple[str, int, int, int]:
    document = Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append("\t".join(cells))
    text = "\n".join(blocks)
    return text, 1, len(blocks), len(document.tables)


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
